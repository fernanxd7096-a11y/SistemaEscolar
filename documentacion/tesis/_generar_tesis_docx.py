from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
import re

doc = Document()
style = doc.styles['Normal']
style.font.name = 'Times New Roman'
style.font.size = Pt(12)


def add_heading(text, level=1):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.name = 'Times New Roman'


def add_para(text, bold=False):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(12)
    run.bold = bold
    return p


title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = title.add_run(
    'Desarrollo de un Sistema Web de Gestión Escolar Integral\n'
    'para la I.E.P. Milagroso San Judas Tadeo'
)
r.bold = True
r.font.size = Pt(14)
r.font.name = 'Times New Roman'

doc.add_paragraph()
for line in [
    'Proyecto de Innovación y/o Mejora',
    'Ingeniería de Software con Inteligencia Artificial',
    'CFP Luis Cáceres Graziani — SENATI',
    '',
    'Autores: [COMPLETAR]',
    'Asesor: Mg. Jose Armando Tiznado Ubillus',
    '',
    'Lima — Perú, 2026',
]:
    p = doc.add_paragraph(line)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_page_break()

with open(r'c:\Users\ferna\Desktop\SistemaEscolar\TESIS_AVANCE_SistemaEscolar.md', encoding='utf-8') as f:
    content = f.read()

sections = re.split(r'\n---\n', content)
skip_first = True
for section in sections:
    if skip_first:
        skip_first = False
        continue
    lines = section.strip().split('\n')
    if not lines or lines[0].startswith('>'):
        continue
    i = 0
    while i < len(lines):
        line = lines[i].strip()
        if not line or line.startswith('*Documento generado'):
            i += 1
            continue
        if line.startswith('# ') and not line.startswith('## '):
            add_heading(line[2:], 1)
        elif line.startswith('## '):
            add_heading(line[3:], 2)
        elif line.startswith('### '):
            add_heading(line[4:], 3)
        elif line.startswith('|') and i + 1 < len(lines) and lines[i + 1].strip().startswith('|'):
            rows = []
            while i < len(lines) and lines[i].strip().startswith('|'):
                row = [c.strip() for c in lines[i].strip().strip('|').split('|')]
                if not all(set(c) <= set('- ') for c in row):
                    rows.append(row)
                i += 1
            if rows:
                table = doc.add_table(rows=len(rows), cols=len(rows[0]))
                table.style = 'Table Grid'
                for ri, row in enumerate(rows):
                    for ci, cell in enumerate(row):
                        table.rows[ri].cells[ci].text = cell
            continue
        elif line.startswith('```'):
            i += 1
            code_lines = []
            while i < len(lines) and not lines[i].strip().startswith('```'):
                code_lines.append(lines[i])
                i += 1
            add_para('\n'.join(code_lines))
        elif line.startswith('- '):
            doc.add_paragraph(line[2:], style='List Bullet')
        elif line.startswith('**') and line.endswith('**'):
            add_para(line.strip('*'), bold=True)
        else:
            add_para(line)
        i += 1
    doc.add_paragraph()

out = r'c:\Users\ferna\Desktop\SistemaEscolar\TESIS_AVANCE_SistemaEscolar.docx'
doc.save(out)
print('Generado:', out)
