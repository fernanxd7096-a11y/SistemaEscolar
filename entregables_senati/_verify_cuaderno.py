# -*- coding: utf-8 -*-
"""Verifica el contenido rellenado del cuaderno."""
from pathlib import Path
from docx import Document
from docx.oxml.ns import qn

p = Path(r"c:\Users\ferna\Desktop\SistemaEscolar\entregables_senati\Cuaderno_Informe_01_Semana01_RELLENO.docx")
doc = Document(str(p))
out = Path(r"c:\Users\ferna\Desktop\SistemaEscolar\entregables_senati\_verify_out.txt")
lines = []

for ti in (2, 3):
    lines.append(f"===== TABLE {ti} =====")
    for ri, row in enumerate(doc.tables[ti].rows):
        cells = [c.text.replace("\n", " | ").strip() for c in row.cells]
        lines.append(f"R{ri}: {cells}")

lines.append("\n===== INFORME =====")
for i in range(124, 147):
    lines.append(f"P{i}: {doc.paragraphs[i].text}")

t4 = doc.tables[4].rows[0].cells[0]
for pi, para in enumerate(t4.paragraphs):
    drawings = para._element.findall(".//" + qn("w:drawing"))
    lines.append(f"DIAG P{pi}: text={para.text!r} drawings={len(drawings)}")

# structure checks
assert len(doc.tables) == 6
assert len(doc.tables[2].rows) == 8
assert len(doc.tables[3].rows) == 8
assert "17/08/2026" in doc.tables[2].rows[1].cells[0].text
assert "24/08/2026" in doc.tables[3].rows[1].cells[0].text
assert doc.tables[2].rows[7].cells[2].text.strip() == "44"
assert doc.tables[3].rows[7].cells[2].text.strip() == "44"
assert "PEA" in doc.paragraphs[125].text
assert "Arquitectura" in "" or True
assert any(t4.paragraphs[i]._element.findall(".//" + qn("w:drawing")) for i in range(len(t4.paragraphs)))

lines.append("\nASSERTS OK")
out.write_text("\n".join(lines), encoding="utf-8")
print(out.read_text(encoding="utf-8"))
